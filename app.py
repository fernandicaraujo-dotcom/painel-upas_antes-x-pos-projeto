"""
Painel Antes x Pós-Projeto — UPAs Porte I e Porte III

Lê a planilha ANALISE_UPAS.xlsx (uma aba por mês) a partir de um link
(Google Sheets, OneDrive/SharePoint, Dropbox ou GitHub) e compara, para cada
UPA, os indicadores do período Antes-Projeto com o Pós-Projeto.

Os dados são recarregados automaticamente a cada 5 minutos (ou pelo botão
"Atualizar dados agora"), então o painel acompanha a planilha conforme ela é
preenchida.
"""

from __future__ import annotations

import datetime as dt
import io
import os
import re
import unicodedata

import numpy as np
import openpyxl
import pandas as pd
import plotly.express as px
import plotly.graph_objects as go
import requests
import streamlit as st

# ----------------------------------------------------------------------------
# Configuração
# ----------------------------------------------------------------------------
st.set_page_config(page_title="Painel UPAs · Antes x Pós-Projeto", page_icon="🏥", layout="wide")

CACHE_SEGUNDOS = 300  # recarrega a planilha a cada 5 minutos
ARQUIVO_LOCAL = os.path.join(os.path.dirname(__file__), "dados", "ANALISE_UPAS.xlsx")

# Paleta IGESDF (extraída da logo): azul #1890C0, ciano #08B8C8, verde-limão #B8D030, amarelo #F8F000, laranja #F06820
IG_AZUL, IG_CIANO, IG_LIMAO, IG_AMARELO, IG_LARANJA = "#1890C0", "#08B8C8", "#B8D030", "#F8F000", "#F06820"
VERDE, VERMELHO, CINZA, AZUL, AZUL_CLARO = "#3E8E1F", "#D2451E", "#8E979F", "#1477A6", "#08A3B2"
ASSETS = os.path.join(os.path.dirname(__file__), "assets")

# Indicadores: chave interna -> (nome exibido, tipo, sentido)
#   tipo:   "volume" (média diária), "tempo" (minutos), "pct" (razão)
#   sentido (definido pelas metas das UPAs no 51º Termo Aditivo ao Contrato 001/2018-SES/DF):
#     +1 maior é melhor  -> atendimentos médicos + urgência (meta mínima) e % de GAEs classificadas
#     -1 menor é melhor  -> tempos médios entre a classificação e o 1º atendimento médico
#      0 sem meta no contrato -> o usuário escolhe na barra lateral (padrão: só variação)
INDICADORES = {
    "producao":    ("Atend. médicos + urgência (meta)", "volume", +1),
    "total":       ("Acolhim. + Urgência + Atend. UPA", "volume", +1),
    "acolhimento": ("Acolhimento c/ Classif. de Risco", "volume", 0),
    "urgencia":    ("Atend. Urgência Obs. ≤ 24h",       "volume", +1),
    "medico":      ("Atend. Médico em UPA",             "volume", +1),
    "exames":      ("Exames Laboratoriais",             "volume", 0),
    "rx":          ("Radiografias",                     "volume", 0),
    "ecg":         ("Eletrocardiogramas",               "volume", 0),
    "gae_ab":      ("GAEs Abertas",                     "volume", 0),
    "gae_cl":      ("GAEs Classificadas",               "volume", 0),
    "gae_pct":     ("% de GAEs classificadas",          "pct",    +1),
    "t_amarelo":   ("Tempo médio · Amarelo",            "tempo",  -1),
    "t_laranja":   ("Tempo médio · Laranja",            "tempo",  -1),
    "t_verde":     ("Tempo médio · Verde",              "tempo",  -1),
}
ORDEM = list(INDICADORES)
NOME = {k: v[0] for k, v in INDICADORES.items()}

MESES = ["", "jan", "fev", "mar", "abr", "mai", "jun", "jul", "ago", "set", "out", "nov", "dez"]


def _sem_acento(s: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn").lower()


def chave_indicador(nome: str) -> str | None:
    n = _sem_acento(str(nome))
    regras = [
        ("acolhim. +", "total"), ("percentual", "gae_pct"), ("amarelo", "t_amarelo"),
        ("laranja", "t_laranja"), ("verde", "t_verde"), ("acolhimento", "acolhimento"),
        ("urgencia", "urgencia"), ("medico", "medico"), ("exames", "exames"),
        ("radiograf", "rx"), ("eletrocard", "ecg"), ("gaes abertas", "gae_ab"),
        ("gaes classificadas", "gae_cl"),
    ]
    for trecho, chave in regras:
        if trecho in n:
            return chave
    return None


def nome_upa(nome) -> str:
    """'UPA\nCEILÂNDIA II' -> 'UPA Ceilândia II'."""
    palavras = re.sub(r"\s+", " ", str(nome)).strip().split(" ")
    saida = []
    for p in palavras:
        if p.upper() in ("UPA", "II", "III", "IV"):
            saida.append(p.upper())
        elif p.lower() in ("de", "da", "das", "do", "dos"):
            saida.append(p.lower())
        else:
            saida.append(p.capitalize())
    return " ".join(saida)


def para_minutos(v) -> float:
    """Converte hh:mm (time, timedelta, datetime, número de dias do Excel ou texto) em minutos."""
    if v is None or v == "":
        return np.nan
    if isinstance(v, dt.timedelta):
        m = v.total_seconds() / 60
    elif isinstance(v, dt.datetime):
        m = v.hour * 60 + v.minute + v.second / 60
    elif isinstance(v, dt.time):
        m = v.hour * 60 + v.minute + v.second / 60
    elif isinstance(v, (int, float)):
        m = float(v) * 24 * 60  # fração de dia do Excel
    else:
        partes = re.findall(r"\d+", str(v))
        if not partes:
            return np.nan
        h, mi, s = (list(map(int, partes)) + [0, 0, 0])[:3]
        m = h * 60 + mi + s / 60
    return m if m > 0 else np.nan  # tempo zerado = sem informação


def para_numero(v) -> float:
    if v is None or (isinstance(v, str) and (v.startswith("=") or not v.strip())):
        return np.nan
    try:
        return float(v)
    except (TypeError, ValueError):
        return np.nan


# ----------------------------------------------------------------------------
# Leitura da planilha
# ----------------------------------------------------------------------------
def link_para_download(url: str) -> str:
    """Transforma links de compartilhamento em links de download direto."""
    url = url.strip()
    g = re.search(r"docs\.google\.com/spreadsheets/d/([\w-]+)", url)
    if g:
        return f"https://docs.google.com/spreadsheets/d/{g.group(1)}/export?format=xlsx"
    gd = re.search(r"drive\.google\.com/(?:file/d/|open\?id=)([\w-]+)", url)
    if gd:
        return f"https://drive.google.com/uc?export=download&id={gd.group(1)}"
    if "github.com" in url and "/blob/" in url:
        return url.replace("github.com", "raw.githubusercontent.com").replace("/blob/", "/")
    if "dropbox.com" in url:
        return re.sub(r"[?&]dl=0", "", url) + ("&" if "?" in url else "?") + "dl=1"
    if any(d in url for d in ("1drv.ms", "onedrive.live.com", "sharepoint.com")) and "download=1" not in url:
        return url + ("&" if "?" in url else "?") + "download=1"
    return url


@st.cache_data(ttl=CACHE_SEGUNDOS, show_spinner="Lendo a planilha…")
def baixar(url: str) -> tuple[bytes, str]:
    r = requests.get(link_para_download(url), timeout=60)
    r.raise_for_status()
    if r.content[:2] != b"PK":
        raise ValueError(
            "O link não devolveu um arquivo Excel. Verifique se a planilha está compartilhada "
            "como 'Qualquer pessoa com o link pode ver'."
        )
    return r.content, dt.datetime.now().strftime("%d/%m/%Y %H:%M")


@st.cache_data(show_spinner=False)
def ler_planilha(conteudo: bytes) -> pd.DataFrame:
    """Devolve uma tabela longa: data, fase, upa, porte, indicador, valor."""
    wb = openpyxl.load_workbook(io.BytesIO(conteudo), data_only=False)
    linhas = []
    for ws in wb.worksheets:
        # Linha de cabeçalho = a que tem "Data" na coluna A
        cab = next((r for r in range(1, 15) if str(ws.cell(r, 1).value or "").strip().lower() == "data"), None)
        if cab is None:
            continue
        upas, porte_atual = {}, ""
        for c in range(4, ws.max_column + 1):
            porte_txt = str(ws.cell(cab - 1, c).value or "") if cab > 1 else ""
            if "PORTE" in porte_txt.upper():
                porte_atual = "Porte " + porte_txt.upper().split("PORTE")[1].strip(" —-")
            nome = ws.cell(cab, c).value
            if nome:
                upas[c] = (nome_upa(nome), porte_atual)
        for r in range(cab + 1, ws.max_row + 1):
            data = ws.cell(r, 1).value
            if not isinstance(data, (dt.datetime, dt.date)):
                continue
            chave = chave_indicador(ws.cell(r, 3).value)
            if chave in (None, "total", "gae_pct"):  # calculados depois, a partir dos dados digitados
                continue
            fase = "Pós" if "pos" in _sem_acento(str(ws.cell(r, 2).value)) else "Antes"
            tipo = INDICADORES[chave][1]
            for c, (upa, porte) in upas.items():
                v = ws.cell(r, c).value
                valor = para_minutos(v) if tipo == "tempo" else para_numero(v)
                linhas.append((pd.Timestamp(data).normalize(), fase, upa, porte, chave, valor))
    df = pd.DataFrame(linhas, columns=["data", "fase", "upa", "porte", "ind", "valor"])
    if df.empty:
        return df
    # Dias ainda não preenchidos: descarta datas sem nenhum valor
    preenchidas = df.groupby("data")["valor"].apply(lambda s: s.notna().any())
    df = df[df["data"].isin(preenchidas[preenchidas].index)]
    # Indicadores calculados
    calc = []
    for chave, partes in (("total", ["acolhimento", "urgencia", "medico"]), ("producao", ["urgencia", "medico"])):
        base = df[df["ind"].isin(partes)]
        t = base.groupby(["data", "fase", "upa", "porte"], as_index=False)["valor"].sum(min_count=1)
        t["ind"] = chave
        calc.append(t)
    return pd.concat([df, *calc], ignore_index=True)


# ----------------------------------------------------------------------------
# Agregação
# ----------------------------------------------------------------------------
def agregar(df: pd.DataFrame, por: list[str]) -> pd.DataFrame:
    """Média diária (volumes), média dos tempos e % de GAEs = soma classificadas / soma abertas."""
    out = (df.groupby(por + ["ind"], as_index=False)["valor"].mean())
    gae = df[df["ind"].isin(["gae_ab", "gae_cl"])].pivot_table(index=por, columns="ind", values="valor", aggfunc="sum")
    if {"gae_ab", "gae_cl"} <= set(gae.columns):
        pct = (gae["gae_cl"] / gae["gae_ab"].replace(0, np.nan) * 100).rename("valor").reset_index()
        pct["ind"] = "gae_pct"
        out = pd.concat([out, pct], ignore_index=True)
    return out


def comparar(df: pd.DataFrame, por: list[str], sentido_volume: int, limiar: float) -> pd.DataFrame:
    ag = agregar(df, por + ["fase"])
    tab = ag.pivot_table(index=por + ["ind"], columns="fase", values="valor").reset_index()
    for f in ("Antes", "Pós"):
        if f not in tab:
            tab[f] = np.nan
    tab["var_abs"] = tab["Pós"] - tab["Antes"]
    tab["var_pct"] = np.where(tab["ind"] == "gae_pct", tab["var_abs"],  # em pontos percentuais
                              tab["var_abs"] / tab["Antes"].replace(0, np.nan) * 100)
    tab["sentido"] = tab["ind"].map(lambda k: INDICADORES[k][2] or sentido_volume)

    def status(row):
        if pd.isna(row["var_pct"]):
            return "Sem dados"
        if row["sentido"] == 0:
            return "Aumentou" if row["var_pct"] > limiar else "Diminuiu" if row["var_pct"] < -limiar else "Estável"
        score = row["var_pct"] * row["sentido"]
        return "Melhorou" if score > limiar else "Piorou" if score < -limiar else "Estável"

    tab["status"] = tab.apply(status, axis=1)
    tab["score"] = np.where(tab["sentido"] == 0, 0, tab["var_pct"] * tab["sentido"])
    tab["indicador"] = tab["ind"].map(NOME)
    tab["ordem"] = tab["ind"].map(ORDEM.index)
    return tab.sort_values(por + ["ordem"])


def fmt_valor(v, ind):
    if pd.isna(v):
        return "—"
    tipo = INDICADORES[ind][1]
    if tipo == "tempo":
        h, m = divmod(int(round(v)), 60)
        return f"{h}h{m:02d}"
    if tipo == "pct":
        return f"{v:.1f}%".replace(".", ",")
    return f"{v:,.1f}".replace(",", "X").replace(".", ",").replace("X", ".")


def fmt_var(v, ind):
    if pd.isna(v):
        return "—"
    unidade = " p.p." if INDICADORES[ind][1] == "pct" else "%"
    return f"{v:+.1f}".replace(".", ",") + unidade


COR_STATUS = {"Melhorou": VERDE, "Piorou": VERMELHO, "Estável": CINZA,
              "Aumentou": AZUL, "Diminuiu": AZUL_CLARO, "Sem dados": "#D0D5DB"}
ICONE = {"Melhorou": "▲ Melhorou", "Piorou": "▼ Piorou", "Estável": "● Estável",
         "Aumentou": "↑ Aumentou", "Diminuiu": "↓ Diminuiu", "Sem dados": "– Sem dados"}


# ----------------------------------------------------------------------------
# Barra lateral: fonte de dados e critérios
# ----------------------------------------------------------------------------
st.sidebar.header("Fonte dos dados")
url_padrao = ""
try:
    url_padrao = st.secrets.get("PLANILHA_URL", "")
except Exception:
    pass
url_padrao = url_padrao or os.environ.get("PLANILHA_URL", "")

url = st.sidebar.text_input("Link da planilha", value=url_padrao,
                            help="Google Sheets, OneDrive/SharePoint, Dropbox ou GitHub. "
                                 "A planilha precisa estar compartilhada para leitura por link.")
arquivo = st.sidebar.file_uploader("…ou envie o arquivo .xlsx", type=["xlsx"])
if st.sidebar.button("🔄 Atualizar dados agora", width="stretch"):
    baixar.clear()
    st.rerun()

conteudo, origem = None, ""
try:
    if arquivo is not None:
        conteudo, origem = arquivo.getvalue(), f"arquivo enviado ({arquivo.name})"
    elif url:
        conteudo, hora = baixar(url)
        origem = f"link configurado · lido em {hora} · atualiza a cada {CACHE_SEGUNDOS // 60} min"
    elif os.path.exists(ARQUIVO_LOCAL):
        with open(ARQUIVO_LOCAL, "rb") as f:
            conteudo = f.read()
        origem = "arquivo dados/ANALISE_UPAS.xlsx do repositório"
except Exception as e:
    st.error(f"Não consegui ler a planilha pelo link. {e}")
    st.stop()

if conteudo is None:
    st.title("Painel UPAs · Antes x Pós-Projeto")
    st.info("Informe o link da planilha na barra lateral (ou envie o arquivo .xlsx) para carregar o painel.")
    st.stop()

dados = ler_planilha(conteudo)
if dados.empty:
    st.error("A planilha não tem o formato esperado (abas mensais com as colunas Data, Fase, Indicador e as UPAs).")
    st.stop()

st.sidebar.header("Período")
d_min, d_max = dados["data"].min().date(), dados["data"].max().date()
periodo_sel = st.sidebar.date_input("Datas analisadas", value=(d_min, d_max), min_value=d_min, max_value=d_max,
                                    format="DD/MM/YYYY",
                                    help="O Antes e o Pós passam a usar só os dias dentro deste intervalo.")
if isinstance(periodo_sel, (list, tuple)) and len(periodo_sel) == 2:
    dados = dados[(dados["data"].dt.date >= periodo_sel[0]) & (dados["data"].dt.date <= periodo_sel[1])]
recorte = st.sidebar.radio("Recorte do Pós", ["Todo", "Últimos 30 dias", "Últimos 7 dias"], horizontal=True,
                           help="Compara só os últimos dias da planilha (todos no Pós) com a base do Antes.")
tipo_dia = st.sidebar.radio("Dias", ["Todos", "Dias úteis", "Finais de semana"], horizontal=True,
                            help="Finais de semana = sábados e domingos. Feriados não são separados.")
n_rec = {"Todo": 0, "Últimos 30 dias": 30, "Últimos 7 dias": 7}[recorte]
if n_rec:
    datas_todas = sorted(dados["data"].unique())
    corte = datas_todas[max(0, len(datas_todas) - n_rec)]
    dados = dados[(dados["fase"] == "Antes") | (dados["data"] >= corte)]
fds = dados["data"].dt.weekday >= 5
if tipo_dia == "Finais de semana":
    dados = dados[fds]
elif tipo_dia == "Dias úteis":
    dados = dados[~fds]
if dados.empty:
    st.warning("Não há dados no período escolhido.")
    st.stop()
if dados["fase"].nunique() < 2:
    st.warning("A seleção atual não inclui dias das duas fases (Antes e Pós). "
               "Escolha um intervalo que comece antes do início do projeto e termine depois.")

st.sidebar.header("UPA")
portes = sorted(dados["porte"].unique())
porte_sel = st.sidebar.multiselect("Porte", portes, default=portes)
dados = dados[dados["porte"].isin(porte_sel)]
if dados.empty:
    st.warning("Selecione ao menos um porte.")
    st.stop()

TODAS = "Todas as UPAs"
upa_sel = st.sidebar.selectbox("Ver indicadores de", [TODAS] + sorted(dados["upa"].unique()),
                               help="Os cards e o gráfico mudam conforme a UPA e o período escolhidos.")
dados_sel = dados if upa_sel == TODAS else dados[dados["upa"] == upa_sel]


st.sidebar.header("Critérios de comparação")
opc_volume = st.sidebar.radio(
    "Indicadores sem meta no contrato (acolhimento, exames, RX, ECG, GAEs abertas/classificadas): aumentar é…",
    ["Só variação (neutro)", "Melhora", "Piora"], index=0,
    help="Pelas metas das UPAs no 51º Termo Aditivo: atendimentos médicos + urgência e % de GAEs "
         "classificadas, aumentar é melhora. Tempos médios por cor, diminuir é melhora.")
sentido_volume = {"Só variação (neutro)": 0, "Melhora": 1, "Piora": -1}[opc_volume]
limiar = st.sidebar.slider("Variação mínima para considerar mudança (%)", 0.0, 20.0, 5.0, 0.5,
                           help="Variações menores que esse valor aparecem como 'Estável'.")
# ----------------------------------------------------------------------------
# Cabeçalho
# ----------------------------------------------------------------------------
def periodo(df, fase):
    d = df.loc[df["fase"] == fase, "data"]
    return (d.min(), d.max(), d.nunique()) if not d.empty else (None, None, 0)

a_ini, a_fim, a_dias = periodo(dados_sel, "Antes")
p_ini, p_fim, p_dias = periodo(dados_sel, "Pós")
f = lambda d: d.strftime("%d/%m/%Y") if d is not None else "—"

logo_a, _, logo_b = st.columns([2, 5, 1.2])


def achar_logo(nome):
    """Procura a logo em assets/ ou solta na raiz do repositório."""
    for pasta in (ASSETS, os.path.dirname(os.path.abspath(__file__))):
        caminho = os.path.join(pasta, nome)
        if os.path.exists(caminho):
            return caminho
    return None


for coluna, nome, largura in ((logo_a, "logo_igesdf.png", 230), (logo_b, "logo_upa24h.png", 130)):
    caminho = achar_logo(nome)
    if caminho:
        coluna.image(caminho, width=largura)
st.markdown(
    f"""<div style="display:grid;grid-template-columns:3fr 2fr 2fr 1fr 2fr;height:6px;border-radius:3px;overflow:hidden;margin:4px 0 8px">
    <div style="background:{IG_AZUL}"></div><div style="background:{IG_CIANO}"></div><div style="background:{IG_LIMAO}"></div>
    <div style="background:{IG_AMARELO}"></div><div style="background:{IG_LARANJA}"></div></div>""", unsafe_allow_html=True)
st.markdown(f"<h1 style='color:{AZUL};margin:0'>Painel UPAs · Antes x Pós-Projeto</h1>", unsafe_allow_html=True)
st.caption(f"**Antes:** {f(a_ini)} a {f(a_fim)} ({a_dias} dias)  ·  **Pós:** {f(p_ini)} a {f(p_fim)} ({p_dias} dias)  ·  Fonte: {origem}")

comp_sel = comparar(dados_sel, [], sentido_volume, limiar)
comp_upa = comparar(dados, ["upa", "porte"], sentido_volume, limiar)
inicio_pos = periodo(dados, "Pós")[0]

titulo_sel = "Todas as UPAs (média da rede)" if upa_sel == TODAS else \
    f"{upa_sel} · {dados_sel['porte'].iloc[0]}"
st.markdown(f"<h2 style='margin:.6rem 0 0'>{titulo_sel}</h2>", unsafe_allow_html=True)

cls = comp_sel[comp_sel["sentido"] != 0]
n_mel, n_pio, n_est = (int((cls["status"] == s).sum()) for s in ("Melhorou", "Piorou", "Estável"))
c1, c2, c3 = st.columns(3)
c1.metric("Indicadores que melhoraram", n_mel)
c2.metric("Indicadores que pioraram", n_pio)
c3.metric("Estáveis", n_est)
st.caption(f"Contagem feita sobre os {len(cls)} indicadores que têm sentido definido "
           "(meta no 51º Termo Aditivo, ou os que você marcou como melhora/piora na barra lateral).")

aba_ind, aba_comp, aba_dados = st.tabs(["Indicadores", "Comparar UPAs", "Tabela completa"])

# ----------------------------------------------------------------------------
# Indicadores: cards agrupados + evolução mensal
# ----------------------------------------------------------------------------
GRUPOS = [
    ("Atendimentos (média por dia)", ["producao", "total", "medico", "urgencia", "acolhimento"]),
    ("Tempo entre a classificação de risco e o 1º atendimento médico", ["t_laranja", "t_amarelo", "t_verde"]),
    ("Classificação de risco (GAEs)", ["gae_pct", "gae_ab", "gae_cl"]),
    ("Exames (média por dia)", ["exames", "rx", "ecg"]),
]


def card(r):
    cor = COR_STATUS[r["status"]]
    return f"""<div style="border:1px solid rgba(128,128,128,.25);border-top:4px solid {cor};border-radius:8px;
        padding:12px 14px;margin-bottom:12px;min-height:150px">
        <div style="font-size:.82rem;opacity:.8">{r['indicador']}</div>
        <div style="font-size:1.6rem;font-weight:700;line-height:1.3">{fmt_valor(r['Pós'], r['ind'])}</div>
        <div style="font-size:.82rem;opacity:.7">antes do projeto: {fmt_valor(r['Antes'], r['ind'])}</div>
        <div style="margin-top:4px;color:{cor};font-weight:700;font-size:.9rem">{ICONE[r['status']]} · {fmt_var(r['var_pct'], r['ind'])}</div>
        </div>"""


with aba_ind:
    st.caption("O número grande é o valor depois do projeto (Pós). Embaixo, o valor de antes e a variação.")
    por_ind = comp_sel.set_index("ind")
    for titulo, chaves in GRUPOS:
        st.markdown(f"**{titulo}**")
        chaves = [k for k in chaves if k in por_ind.index]
        cols = st.columns(len(chaves) if len(chaves) > 3 else 3)
        for col, k in zip(cols, chaves):
            r = por_ind.loc[k].copy()
            r["ind"] = k
            col.markdown(card(r), unsafe_allow_html=True)

    st.markdown("**Evolução mês a mês**")
    ind_sel = st.selectbox("Indicador", ORDEM, format_func=NOME.get, key="ind_evol")
    com_mes = lambda df: df.assign(mes=df["data"].dt.to_period("M").dt.to_timestamp())
    serie = agregar(com_mes(dados_sel), ["mes"])
    serie = serie[serie["ind"] == ind_sel].sort_values("mes")
    fig = go.Figure()
    if upa_sel != TODAS:
        rede_m = agregar(com_mes(dados), ["mes"])
        rede_m = rede_m[rede_m["ind"] == ind_sel].sort_values("mes")
        fig.add_scatter(x=rede_m["mes"], y=rede_m["valor"], name="Média de todas as UPAs", mode="lines",
                        line=dict(color=CINZA, dash="dot"))
    fig.add_scatter(x=serie["mes"], y=serie["valor"], name=upa_sel if upa_sel != TODAS else "Média de todas as UPAs",
                    mode="lines+markers+text", line=dict(color=AZUL, width=3),
                    text=[fmt_valor(v, ind_sel) for v in serie["valor"]], textposition="top center")
    if inicio_pos is not None:
        fig.add_vline(x=inicio_pos, line_dash="dash", line_color=IG_LARANJA)
        fig.add_annotation(x=inicio_pos, y=1.05, yref="paper", text="Início do projeto", showarrow=False,
                           xanchor="left", font=dict(color=IG_LARANJA))
    unid = {"tempo": "minutos", "pct": "%", "volume": "média por dia"}[INDICADORES[ind_sel][1]]
    fig.update_layout(height=380, margin=dict(l=10, r=10, t=40, b=10), yaxis_title=unid,
                      legend=dict(orientation="h", y=-0.15), xaxis=dict(tickformat="%b/%y"))
    st.plotly_chart(fig, width="stretch")
    st.caption("Cada ponto é a média do mês. A linha laranja marca o início do projeto (junho mistura as duas fases).")

# ----------------------------------------------------------------------------
# Comparar UPAs
# ----------------------------------------------------------------------------
with aba_comp:
    st.markdown("**Resumo por UPA**")
    st.caption("Para cada UPA: quantos indicadores melhoraram e pioraram, e qual foi a maior melhora e a maior piora.")
    linhas = []
    for (upa, porte), t in comp_upa.groupby(["upa", "porte"]):
        tc = t[(t["sentido"] != 0) & t["score"].notna()]
        mel, pio = tc[tc["status"] == "Melhorou"], tc[tc["status"] == "Piorou"]
        top = mel.loc[mel["score"].idxmax()] if not mel.empty else None
        pior = pio.loc[pio["score"].idxmin()] if not pio.empty else None
        linhas.append({
            "UPA": upa, "Porte": porte,
            "Melhoraram": len(mel), "Pioraram": len(pio), "Estáveis": int((tc["status"] == "Estável").sum()),
            "Maior melhora": f"{top['indicador']} ({fmt_var(top['var_pct'], top['ind'])})" if top is not None else "—",
            "Maior piora": f"{pior['indicador']} ({fmt_var(pior['var_pct'], pior['ind'])})" if pior is not None else "—",
        })
    resumo = pd.DataFrame(linhas)
    resumo = resumo.assign(_s=resumo["Melhoraram"] - resumo["Pioraram"]).sort_values(["_s", "Melhoraram"], ascending=False).drop(columns="_s")

    def cores(col):
        if col.name == "Melhoraram":
            return [f"color:{VERDE};font-weight:700"] * len(col)
        if col.name == "Pioraram":
            return [f"color:{VERMELHO};font-weight:700"] * len(col)
        return [""] * len(col)

    st.dataframe(resumo.style.apply(cores), hide_index=True, width="stretch", height=(len(resumo) + 1) * 35 + 3)

    st.markdown("**Um indicador em todas as UPAs**")
    ind = st.selectbox("Indicador", ORDEM, format_func=NOME.get, key="ind_comp")
    t = comp_upa[comp_upa["ind"] == ind].sort_values("Pós")
    fig = go.Figure()
    fig.add_bar(y=t["upa"], x=t["Antes"], name="Antes do projeto", orientation="h", marker_color="#B8C2CE",
                text=[fmt_valor(v, ind) for v in t["Antes"]], textposition="outside")
    fig.add_bar(y=t["upa"], x=t["Pós"], name="Depois do projeto", orientation="h",
                marker_color=[COR_STATUS[s] if s in ("Melhorou", "Piorou") else AZUL for s in t["status"]],
                text=[f"{fmt_valor(v, ind)}  ({fmt_var(p, ind)})" for v, p in zip(t["Pós"], t["var_pct"])],
                textposition="outside")
    fig.update_layout(barmode="group", height=80 + 52 * len(t), margin=dict(l=10, r=60, t=10, b=10),
                      legend=dict(orientation="h", y=1.04))
    st.plotly_chart(fig, width="stretch")
    st.caption("Barra cinza = antes do projeto. Barra colorida = depois: verde se melhorou, vermelha se piorou, azul se ficou estável ou não tem meta.")

# ----------------------------------------------------------------------------
# Tabela completa
# ----------------------------------------------------------------------------
with aba_dados:
    saida = comp_upa[["upa", "porte", "indicador", "Antes", "Pós", "var_pct", "status"]].rename(columns={
        "upa": "UPA", "porte": "Porte", "indicador": "Indicador", "var_pct": "Variação (%/p.p.)", "status": "Situação"})
    st.dataframe(saida.round(2), hide_index=True, width="stretch", height=520)
    st.download_button("Baixar tabela (CSV)", saida.round(2).to_csv(index=False, sep=";", decimal=",").encode("utf-8-sig"),
                       "comparativo_upas.csv", "text/csv")
    st.caption("Tempos em minutos. Volumes em média por dia. % de GAEs com variação em pontos percentuais.")
